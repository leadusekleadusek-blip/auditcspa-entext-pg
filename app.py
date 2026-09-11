import base64
import json
import io
import os
import requests
from datetime import datetime
import streamlit as st
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Configuration de la page
st.set_page_config(
    page_title="Audit Sécurité & HSE - Entreprises Extérieures",
    page_icon="🛡️",
    layout="wide",
    initial_sidebar_state="expanded"
)

# Clés Airtable depuis les secrets Streamlit
AIRTABLE_TOKEN = st.secrets.get("AIRTABLE_TOKEN", "")
AIRTABLE_BASE_ID = st.secrets.get("AIRTABLE_BASE_ID", "")
AIRTABLE_TABLE_NAME = st.secrets.get("AIRTABLE_TABLE_NAME", "Audits")

DB_FILE = "audits_db.json"

if "local_audits" not in st.session_state:
    st.session_state["local_audits"] = []

# Style CSS
st.markdown("""
    <style>
    .main-header { font-size: 28px; font-weight: bold; color: #1E3A8A; margin-bottom: 5px; }
    .sub-header { font-size: 15px; color: #4B5563; margin-bottom: 20px; }
    
    .question-card {
        background-color: #FFFFFF;
        border: 2px solid #CBD5E1;
        border-radius: 10px;
        padding: 20px;
        margin-bottom: 15px;
        box-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.1), 0 2px 4px -1px rgba(0, 0, 0, 0.06);
    }
    
    .question-title-big {
        font-size: 20px !important;
        font-weight: 700 !important;
        color: #1E3A8A !important;
        margin-bottom: 5px;
        border-bottom: 2px solid #E2E8F0;
        padding-bottom: 6px;
    }
    
    .guidance-box { 
        background-color: #F0F9FF; 
        border-left: 4px solid #0284C7; 
        padding: 12px 16px; 
        border-radius: 6px; 
        font-size: 14px; 
        color: #0369A1; 
        margin-top: 8px; 
        margin-bottom: 16px; 
    }
    
    .stButton > button {
        width: 100% !important;
        font-size: 20px !important;
        font-weight: bold !important;
        padding: 16px 28px !important;
        background-color: #1E3A8A !important;
        color: #FFFFFF !important;
        border: 3px solid #3B82F6 !important;
        border-radius: 10px !important;
        cursor: pointer !important;
    }
    .stButton > button:hover {
        background-color: #2563EB !important;
        border-color: #60A5FA !important;
        color: #FFFFFF !important;
    }
    </style>
""", unsafe_allow_html=True)

QUESTIONS_DATA = [
    # Thème 1 : Attentes & Engagement
    {"id": "1.1", "theme": "1. Attentes & Engagement", "cat": "1. Attentes & Engagement", "q": "Programme de sécurité complet en place sur le site", "g": "• Manuel numérique spécifique au site rédigé et documenté\n• Doit englober tous les aspects des livrables sécurité définis dans le cahier des charges"},
    {"id": "1.2", "theme": "1. Attentes & Engagement", "cat": "1. Attentes & Engagement", "q": "Plan d'action d'urgence spécifique au site", "g": "• Plan tenu à jour incluant la chaîne d'alerte et les responsabilités de chacun"},
    {"id": "1.3", "theme": "1. Attentes & Engagement", "cat": "1. Attentes & Engagement", "q": "Processus d'analyse et d'enquête sur les incidents", "g": "• Procédure conforme aux exigences HSE du site"},
    {"id": "1.4", "theme": "1. Attentes & Engagement", "cat": "1. Attentes & Engagement", "q": "Implication du personnel à tous les niveaux dans la sécurité", "g": "• Échanger avec le personnel de terrain pour s'assurer que leurs besoins sont pris en compte"},
    {"id": "1.5", "theme": "1. Attentes & Engagement", "cat": "1. Attentes & Engagement", "q": "Qualification des sous-traitants et des prestations gérées", "g": "• Système de qualification préalable des sous-traitants\n• Procédure d'accueil pour tout nouveau prestataire"},

    # Thème 2 : Vision & Suivi de la Performance
    {"id": "2.1", "theme": "2. Vision & Suivi de la Performance", "cat": "2. Vision & Suivi de la Performance", "q": "Vision et objectifs sécurité spécifiques au site (revus annuellement)", "g": "• Objectif principal : zéro accident\n• Objectifs communiqués à l'ensemble du personnel"},
    {"id": "2.2", "theme": "2. Vision & Suivi de la Performance", "cat": "2. Vision & Suivi de la Performance", "q": "Engagement dans l'amélioration continue de la culture sécurité", "g": "• Partage des retours d'expérience et déploiement de nouvelles initiatives sécurité"},
    {"id": "2.3", "theme": "2. Vision & Suivi de la Performance", "cat": "2. Vision & Suivi de la Performance", "q": "Réunions mensuelles minimum entre le responsable sécurité prestataire et le HSE site", "g": "• Garantir des objectifs communs et une communication fluide"},
    {"id": "2.4", "theme": "2. Vision & Suivi de la Performance", "cat": "2. Vision & Suivi de la Performance", "q": "Création d'objectifs à court terme basés sur les indicateurs terrain", "g": "• Mettre en place des actions préventives basées sur les observations et presqu'accidents"},
    {"id": "2.5", "theme": "2. Vision & Suivi de la Performance", "cat": "2. Vision & Suivi de la Performance", "q": "Enquêtes de perception de la sécurité réalisées", "g": "• Évaluation de l'efficacité du programme sécurité (réalisée tous les 2 ans)"},
    {"id": "2.6", "theme": "2. Vision & Suivi de la Performance", "cat": "2. Vision & Suivi de la Performance", "q": "Suivi, affichage et transmission des indicateurs sécurité au service HSE", "g": "• Heures travaillées, accidents, presqu'accidents, soins infirmiers, visites d'observation"},
    {"id": "2.7", "theme": "2. Vision & Suivi de la Performance", "cat": "2. Vision & Suivi de la Performance", "q": "Utilisation d'outils numériques et tableaux de bord de suivi", "g": "• Numérisation du suivi sécurité pour un accès immédiat aux données"},
    {"id": "2.8", "theme": "2. Vision & Suivi de la Performance", "cat": "2. Vision & Suivi de la Performance", "q": "Revues de performance sécurité annuelles", "g": "• Bilan annuel du personnel et partage des retours d'expérience"},

    # Thème 3 : Risques Critiques & Meilleures Pratiques
    {"id": "3.1.1", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.1 Risques Vitaux (Life Critical)", "q": "Intervention en Espaces Confinés", "g": "Formation du personnel, permis de pénétrer, surveillant qualifié et moyens de secours"},
    {"id": "3.1.2", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.1 Risques Vitaux (Life Critical)", "q": "Opérations de Levage et Grutage", "g": "Conducteurs certifiés (CACES), contrôle quotidien des engins, plan de levage si nécessaire"},
    {"id": "3.1.3", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.1 Risques Vitaux (Life Critical)", "q": "Travaux Électriques et Habilitations", "g": "Personnel habilité, travaux sous tension exceptionnels et encadrés, EPI contre le risque d'arc électrique"},
    {"id": "3.1.4", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.1 Risques Vitaux (Life Critical)", "q": "Fouilles, Excavations et Tranchées", "g": "Repérage des réseaux enterrés, personne compétente sur site, permis de fouille"},
    {"id": "3.1.5", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.1 Risques Vitaux (Life Critical)", "q": "Prévention des Chutes de Hauteur", "g": "Port du harnais obligatoire dès 1,80 m et en nacelle, plan de prévention antichute écrit"},
    {"id": "3.1.6", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.1 Risques Vitaux (Life Critical)", "q": "Gestion des Produits Chimiques", "g": "Fiches de Données de Sécurité (FDS) disponibles, étiquetage réglementaire des produits"},
    {"id": "3.1.7", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.1 Risques Vitaux (Life Critical)", "q": "Consignation et Condamnation d'Énergies (LOTO)", "g": "Procédure LOTO formalisée, condamnation par cadenas individuel, traçabilité des opérations"},
    {"id": "3.1.8", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.1 Risques Vitaux (Life Critical)", "q": "Montage et Utilisation d'Échafaudages", "g": "Montage par du personnel qualifié, vérification quotidienne et affichage des panneaux de conformité"},
    {"id": "3.1.9", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.1 Risques Vitaux (Life Critical)", "q": "Utilisation Sécurisée des Échelles", "g": "Échelles industrielles isolantes, usage limité aux accès ou travaux ponctuels très légers"},
    {"id": "3.1.10", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.1 Risques Vitaux (Life Critical)", "q": "Travaux sur Toitures", "g": "Accès sécurisés, balisage et garde-corps ou lignes de vie en place"},
    {"id": "3.1.11", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.1 Risques Vitaux (Life Critical)", "q": "Nacelles et Plateformes Élévatrices (PEMPS)", "g": "Conducteurs formés et titulaires de l'autorisation de conduite (CACES)"},
    {"id": "3.1.12", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.1 Risques Vitaux (Life Critical)", "q": "Élingage et Matériel de Levage", "g": "Inspection périodique des élingues et accessoires, élingueurs qualifiés"},
    {"id": "3.2.1", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.2 Meilleures Pratiques", "q": "Validation Préalable des Produits Chimiques", "g": "Liste des produits autorisés sur site, interdiction des produits non validés au préalable"},
    {"id": "3.2.2", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.2 Meilleures Pratiques", "q": "Contrôle des Expositions Chimiques et Poussières", "g": "Mesures de prévention (extraction, ventilation) et port d'EPI adaptés (chrome, poussières)"},
    {"id": "3.2.3", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.2 Meilleures Pratiques", "q": "Signalement des Incidents et Presqu'accidents", "g": "Enregistrement rapide de tous les événements sécurité dans la base de données"},
    {"id": "3.2.4", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.2 Meilleures Pratiques", "q": "Travaux de Démolition", "g": "Plan de démolition ciblé (gestion de l'amiante, du plomb, stabilité des structures)"},
    {"id": "3.2.5", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.2 Meilleures Pratiques", "q": "Gestion des Situations d'Urgence", "g": "Consignes d'urgence, numéros utiles et plans d'évacuation connus des intervenants"},
    {"id": "3.2.6", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.2 Meilleures Pratiques", "q": "Gestion de l'Environnement et des Déchets", "g": "Stockage conforme des déchets dangereux, kits anti-pollution disponibles"},
    {"id": "3.2.7", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.2 Meilleures Pratiques", "q": "Ergonomie et Manutention Manuelle", "g": "Matériel d'aide à la manutention, sensibilisation aux gestes et postures pour éviter les TMS"},
    {"id": "3.2.8", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.2 Meilleures Pratiques", "q": "Stockage et Manipulation des Bouteilles de Gaz", "g": "Stockage vertical, bouteilles arrimées et chapeaux de protection en place"},
    {"id": "3.2.9", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.2 Meilleures Pratiques", "q": "Inspections des Organismes Officiels (Inspection du Travail, etc.)", "g": "Procédure d'accueil et d'accompagnement définie en cas de contrôle externe"},
    {"id": "3.2.10", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.2 Meilleures Pratiques", "q": "Prévention des Blessures aux Mains", "g": "Port de gants adaptés au risque, utilisation de cutters de sécurité"},
    {"id": "3.2.11", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.2 Meilleures Pratiques", "q": "Systèmes et Équipements à Risques", "g": "Inventaire des équipements à risques, délivrance de permis de travail spécifiques"},
    {"id": "3.2.12", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.2 Meilleures Pratiques", "q": "Prévention des Incendies (Permis de Feu)", "g": "Délivrance systématique d'un permis de feu pour points chauds et présence d'un piquet d'incendie"},
    {"id": "3.2.13", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.2 Meilleures Pratiques", "q": "Protection contre le Bruit", "g": "Signalisation des zones bruyantes et port obligatoire des protections auditives"},
    {"id": "3.2.14", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.2 Meilleures Pratiques", "q": "Tenue et Propreté du Chantier", "g": "Rangement quotidien, suppression des risques de trébuchement, voies de circulation dégagées"},
    {"id": "3.2.15", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.2 Meilleures Pratiques", "q": "Sécurité des Rayonnements Laser", "g": "Mise en œuvre sécurisée des outils laser et balisage de la zone d'exposition"},
    {"id": "3.2.16", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.2 Meilleures Pratiques", "q": "Prévention de l'Exposition au Plomb", "g": "Repérage préalable, information des salariés et retrait par du personnel qualifié"},
    {"id": "3.2.17", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.2 Meilleures Pratiques", "q": "Ouverture de Tuyauteries et Lignes (Line Breaking)", "g": "Permis spécifique, vidange préalable et port des EPI adaptés"},
    {"id": "3.2.18", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.2 Meilleures Pratiques", "q": "Premiers Secours et Soins d'Urgence", "g": "Trousses de secours approvisionnées, secouristes du travail (SST) identifiés"},
    {"id": "3.2.19", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.2 Meilleures Pratiques", "q": "Équipements de Protection Individuelle (EPI)", "g": "EPI adaptés fournis, port effectif vérifié et sensibilisation du personnel"},
    {"id": "3.2.20", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.2 Meilleures Pratiques", "q": "Sécurité des Rayonnements Ionisants", "g": "Repérage des sources et coordination avec le PCR (Conseiller en Radioprotection)"},
    {"id": "3.2.21", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.2 Meilleures Pratiques", "q": "Sécurité Ferroviaire sur Site", "g": "Respect des gabarits, consignation des voies et signalisation spécifique"},
    {"id": "3.2.22", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.2 Meilleures Pratiques", "q": "Valorisation des Comportements Sûrs", "g": "Programme de reconnaissance encourageant les bonnes pratiques sécurité"},
    {"id": "3.2.23", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.2 Meilleures Pratiques", "q": "Protection Respiratoire", "g": "Choix des masques adaptés, tests d'ajustement (fit-test) et règles de port"},
    {"id": "3.2.24", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.2 Meilleures Pratiques", "q": "Signalisation et Balisage des Dangers", "g": "Panneaux d'avertissement et barrières physiques installés en face du danger"},
    {"id": "3.2.25", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.2 Meilleures Pratiques", "q": "Prévention des Addictions (Alcool / Drogues)", "g": "Politique d'interdiction formalisée et contrôles applicables aux sous-traitants"},
    {"id": "3.2.26", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.2 Meilleures Pratiques", "q": "Conduite de Chariots et Engins de Chantier", "g": "Autorisations de conduite à jour et vérification des compétences des opérateurs"},
    {"id": "3.2.27", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.2 Meilleures Pratiques", "q": "Outillage Électroportatif", "g": "Outillage conforme, inspecté avant emploi et personnel formé"},
    {"id": "3.2.28", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.2 Meilleures Pratiques", "q": "Risques Biologiques et Pathogènes", "g": "Mesures de prévention contre les risques d'exposition biologique"},
    {"id": "3.2.29", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.2 Meilleures Pratiques", "q": "Gestion des Non-Conformités Sécurité", "g": "Procédure disciplinaire pouvant aller jusqu'à l'exclusion du site en cas de manquement grave"},
    {"id": "3.2.30", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.2 Meilleures Pratiques", "q": "Coactivité près des Zones en Exploitation", "g": "Balisage, isolement des zones de travail et coordination avec l'exploitation"},
    {"id": "3.2.31", "theme": "3. Risques Critiques & Meilleures Pratiques", "cat": "3.2 Meilleures Pratiques", "q": "Outillage à Main Non Motorisé", "g": "Choix d'outils sécurisés (ex: cutters rétractables) et état général contrôlé"},

    # Thème 4 : Planification & Opérations
    {"id": "4.1", "theme": "4. Planification & Opérations", "cat": "4. Planification & Opérations", "q": "Utilisation d'outils d'analyse des risques au poste", "g": "• Analyses de tâches (AST / JSA), évaluations préalables (PPSA) et causeries quotidiennes"},
    {"id": "4.2", "theme": "4. Planification & Opérations", "cat": "4. Planification & Opérations", "q": "Visites d'Observation Sécurité (VOS / BOS)", "g": "• Réalisation de visites d'observation à tous les niveaux et suivi des actions associées"},
    {"id": "4.3", "theme": "4. Planification & Opérations", "cat": "4. Planification & Opérations", "q": "Gestion de la Fatigue et des Heures Supplémentaires", "g": "• Planification du temps de travail conforme et gestion de l'aptitude au poste"},
    {"id": "4.4", "theme": "4. Planification & Opérations", "cat": "4. Planification & Opérations", "q": "Contrôles Réguliers des Équipements", "g": "• Suivi des contrôles périodiques (engins, nacelles, extincteurs, harnais, outillage)"},
    {"id": "4.5", "theme": "4. Planification & Opérations", "cat": "4. Planification & Opérations", "q": "Rondes et Visites de Chantier par le Management", "g": "• Contrôles réguliers axés sur le rangement, les accès, l'éclairage et la coactivité"},
    {"id": "4.6", "theme": "4. Planification & Opérations", "cat": "4. Planification & Opérations", "q": "Réunions et Échanges Sécurité", "g": "• Réunions sécurité quotidiennes/hebdomadaires, quarts d'heure sécurité et revues d'incidents"},
    {"id": "4.7", "theme": "4. Planification & Opérations", "cat": "4. Planification & Opérations", "q": "Encadrement HSE du Chantier", "g": "• Dimensionnement adapté des responsables et préventeurs sécurité sur le terrain"},

    # Thème 5 : Formations & Habilitations Sécurité
    {"id": "5.1", "theme": "5. Formations & Habilitations Sécurité", "cat": "5. Formations & Habilitations Sécurité", "q": "Accueil Sécurité Spécifique au Site", "g": "• Accueil complet couvrant les risques du site, les consignes d'urgence et validation des acquis"},
    {"id": "5.2", "theme": "5. Formations & Habilitations Sécurité", "cat": "5. Formations & Habilitations Sécurité", "q": "Registres des Formations et Habilitations du Personnel", "g": "• Registre à jour : SST, CACES, habilitations électriques, travaux en hauteur, etc."},
    {"id": "5.3", "theme": "5. Formations & Habilitations Sécurité", "cat": "5. Formations & Habilitations Sécurité", "q": "Causeries et Quarts d'Heure Sécurité Hebdomadaires", "g": "• Organisation régulière de réunions d'information et d'échanges sécurité avec les équipes"},
    {"id": "5.4", "theme": "5. Formations & Habilitations Sécurité", "cat": "5. Formations & Habilitations Sécurité", "q": "Formation des Nouveaux Arrivants et des Encadrants", "g": "• Formations à la réalisation des visites d'observation et compétences requises"},
    {"id": "5.5", "theme": "5. Formations & Habilitations Sécurité", "cat": "5. Formations & Habilitations Sécurité", "q": "Suivi des Recyclages de Formation", "g": "• Système d'alerte pour le renouvellement à jour des habilitations et recyclages"}
]

THEME_LIST = [
    "1. Attentes & Engagement",
    "2. Vision & Suivi de la Performance",
    "3. Risques Critiques & Meilleures Pratiques",
    "4. Planification & Opérations",
    "5. Formations & Habilitations Sécurité"
]

def enregistrer_audit_fichier_local(record):
    """Sauvegarde locale d'urgence."""
    audits = []
    if os.path.exists(DB_FILE):
        try:
            with open(DB_FILE, "r", encoding="utf-8") as f:
                audits = json.load(f)
        except Exception:
            audits = []
    audits.append(record)
    with open(DB_FILE, "w", encoding="utf-8") as f:
        json.dump(audits, f, ensure_ascii=False, indent=2)

def charger_audits_airtable():
    """Charge les audits depuis Airtable."""
    if not AIRTABLE_TOKEN or not AIRTABLE_BASE_ID:
        return []
    
    url = f"https://api.airtable.com/v0/{AIRTABLE_BASE_ID}/{AIRTABLE_TABLE_NAME}"
    headers = {"Authorization": f"Bearer {AIRTABLE_TOKEN}"}
    records = []
    offset = None

    try:
        while True:
            params = {"offset": offset} if offset else {}
            resp = requests.get(url, headers=headers, params=params, timeout=10)
            if resp.status_code == 200:
                data = resp.json()
                for r in data.get("records", []):
                    fields = r.get("fields", {})
                    if "Audit_Data" in fields:
                        try:
                            records.append(json.loads(fields["Audit_Data"]))
                        except Exception:
                            records.append(fields)
                    else:
                        records.append(fields)
                offset = data.get("offset")
                if not offset:
                    break
            else:
                break
    except Exception:
        pass
    return records

def enregistrer_audit_airtable(record):
    """Envoie une nouvelle ligne d'audit vers Airtable."""
    if not AIRTABLE_TOKEN or not AIRTABLE_BASE_ID:
        return False
    
    url = f"https://api.airtable.com/v0/{AIRTABLE_BASE_ID}/{AIRTABLE_TABLE_NAME}"
    headers = {
        "Authorization": f"Bearer {AIRTABLE_TOKEN}",
        "Content-Type": "application/json"
    }
    
    payload = {
        "fields": {
            "Entreprise": str(record.get("Entreprise", "")),
            "Déclarant": str(record.get("Déclarant", "")),
            "Site": str(record.get("Site", "")),
            "Date": str(record.get("Date", "")),
            "Audit_Data": json.dumps(record, ensure_ascii=False)
        }
    }
    
    try:
        resp = requests.post(url, headers=headers, json=payload, timeout=10)
        return resp.status_code in (200, 201)
    except Exception:
        return False

def get_color_badge(percentage):
    if percentage >= 80:
        return f"🟢 **{percentage:.1f}% (Conforme / Excellent)**"
    elif percentage >= 50:
        return f"🟠 **{percentage:.1f}% (À améliorer)**"
    else:
        return f"🔴 **{percentage:.1f}% (Non conforme / Risque élevé)**"

def generer_excel_formatted(selected_data):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rapport Audit"

    data_dict = selected_data.to_dict() if isinstance(selected_data, pd.Series) else dict(selected_data)

    col_widths = {'A': 14, 'B': 25, 'C': 45, 'D': 16, 'E': 45}
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    ws.merge_cells('A1:E1')
    banner = ws['A1']
    banner.value = "INFORMATIONS DE L'ENTREPRISE EXTÉRIEURE"
    banner.font = Font(name='Calibri', size=13, bold=True, color="FFFFFF")
    banner.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    banner.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 28

    oui_c = sum(1 for col, val in data_dict.items() if str(col).endswith("_Réponse") and str(val).strip() == "Oui")
    non_c = sum(1 for col, val in data_dict.items() if str(col).endswith("_Réponse") and str(val).strip() == "Non")
    tot_app = oui_c + non_c
    score_p = (oui_c / tot_app * 100) if tot_app > 0 else 0.0

    label_font = Font(name='Calibri', bold=True, color="1F4E78")

    ws['A3'] = "Entreprise :"
    ws['A3'].font = label_font
    ws['B3'] = str(data_dict.get("Entreprise", "N/A"))

    ws['D3'] = "Date de l'audit :"
    ws['D3'].font = label_font
    ws['E3'] = str(data_dict.get("Date", "N/A"))

    ws['A4'] = "Site :"
    ws['A4'].font = label_font
    ws['B4'] = str(data_dict.get("Site", "N/A"))

    ws['D4'] = "Déclarant / Auditeur :"
    ws['D4'].font = label_font
    ws['E4'] = str(data_dict.get("Déclarant", "N/A"))

    ws['A5'] = "Score de conformité :"
    ws['A5'].font = label_font
    ws['B5'] = f"{score_p:.1f}%"

    headers = ["N° question", "Thème", "Question", "Réponse", "Justification"]
    ws.row_dimensions[7].height = 24

    header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    header_font = Font(name='Calibri', size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for col_idx, text in enumerate(headers, start=1):
        cell = ws.cell(row=7, column=col_idx, value=text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border

    for row_idx, q in enumerate(QUESTIONS_DATA, start=8):
        q_id = q["id"]
        reponse_val = str(data_dict.get(f"{q_id}_Réponse", "")).strip()
        justif_val = str(data_dict.get(f"{q_id}_Justification", "")).strip()

        row_vals = [q_id, q["theme"], q["q"], reponse_val, justif_val]
        ws.row_dimensions[row_idx].height = 22

        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border
            is_center = col_idx in (1, 4)
            cell.alignment = Alignment(
                horizontal='center' if is_center else 'left',
                vertical='center',
                wrap_text=True
            )

            if col_idx == 4:
                if reponse_val in ["Oui", "Conforme"]:
                    cell.font = Font(name='Calibri', bold=True, color="16A34A")
                elif reponse_val in ["Non", "Non conforme"]:
                    cell.font = Font(name='Calibri', bold=True, color="DC2626")

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

def charger_tous_les_audits():
    """Charge depuis Airtable en priorité, puis fusionne avec le local."""
    audits_list = charger_audits_airtable()
    
    if os.path.exists(DB_FILE):
        try:
            with open(DB_FILE, "r", encoding="utf-8") as f:
                local_data = json.load(f)
                for rec in local_data:
                    if rec not in audits_list:
                        audits_list.append(rec)
        except Exception:
            pass
            
    if "local_audits" in st.session_state and st.session_state["local_audits"]:
        for record in st.session_state["local_audits"]:
            if record not in audits_list:
                audits_list.append(record)

    if audits_list:
        return pd.DataFrame(audits_list)
    return pd.DataFrame()

# ==========================================
# SELECTION DU MODE DE NAVIGATION (SIDEBAR)
# ==========================================
st.sidebar.title("📌 Menu Navigation")
app_mode = st.sidebar.radio(
    "Choisir l'espace :",
    options=["📝 Formulaire Prestataire", "🔒 Espace Administrateur HSE"],
    key="navigation_mode"
)
st.sidebar.markdown("---")

# ==========================================
# PANNEAU LATÉRAL DYNAMIQUE SELON LE MODE
# ==========================================
if app_mode == "📝 Formulaire Prestataire":
    total_q = len(QUESTIONS_DATA)
    answered_q_count = 0
    total_oui = 0
    total_non = 0
    total_na = 0

    theme_stats = {t: {"oui": 0, "non": 0, "na": 0, "answered": 0, "total": 0} for t in THEME_LIST}

    for q in QUESTIONS_DATA:
        theme_stats[q["theme"]]["total"] += 1
        q_key = f"status_{q['id']}"
        if q_key in st.session_state and st.session_state[q_key] is not None:
            answered_q_count += 1
            val = st.session_state[q_key]
            theme_stats[q["theme"]]["answered"] += 1
            if val == "Oui":
                total_oui += 1
                theme_stats[q["theme"]]["oui"] += 1
            elif val == "Non":
                total_non += 1
                theme_stats[q["theme"]]["non"] += 1
            elif val == "N/A":
                total_na += 1
                theme_stats[q["theme"]]["na"] += 1

    total_applicable = total_oui + total_non
    global_score_pct = (total_oui / total_applicable * 100.0) if total_applicable > 0 else 0.0

    st.sidebar.markdown("### 🏆 Score Global en Direct")
    st.sidebar.metric("Conformité Globale", f"{global_score_pct:.1f} %")
    st.sidebar.progress(global_score_pct / 100.0)
    st.sidebar.markdown(f"**Statut :** {get_color_badge(global_score_pct)}")

    progress_ratio = answered_q_count / total_q
    st.sidebar.markdown(f"**Remplissage :** {answered_q_count} / {total_q} questions ({int(progress_ratio*100)}%)")
    st.sidebar.progress(progress_ratio)

    st.sidebar.markdown("---")
    st.sidebar.markdown("### 📌 Navigation par Thème")

    for idx, t_name in enumerate(THEME_LIST):
        ts = theme_stats[t_name]
        t_app = ts["oui"] + ts["non"]
        t_score = (ts["oui"] / t_app * 100.0) if t_app > 0 else 0.0
        
        status_icon = "🟢" if t_score >= 80 else ("🟠" if t_score >= 50 else "🔴")
        st.sidebar.markdown(f"**{status_icon} Thème {idx+1} : {t_score:.0f}%**")
        st.sidebar.caption(f"*{t_name.split('. ')[1]}* — ({ts['answered']}/{ts['total']} renseig.)")

    selected_theme_nav = st.sidebar.radio(
        "🎯 Afficher un thème spécifique :",
        options=["📋 Tous les thèmes"] + THEME_LIST,
        key="nav_theme_selection"
    )

else:
    st.sidebar.markdown("### 📊 Traçabilité Administrateur")
    df_admin_side = charger_tous_les_audits()

    if not df_admin_side.empty:
        total_audits_count = len(df_admin_side)
        scores_list = []
        
        admin_options_sidebar = []
        for idx, row in df_admin_side.iterrows():
            o_c = sum(1 for c in row.index if str(c).endswith("_Réponse") and str(row[c]).strip() == "Oui")
            n_c = sum(1 for c in row.index if str(c).endswith("_Réponse") and str(row[c]).strip() == "Non")
            t_app = o_c + n_c
            sc = (o_c / t_app * 100) if t_app > 0 else 0.0
            scores_list.append(sc)
            
            badge_icon = "🟢" if sc >= 80 else ("🟠" if sc >= 50 else "🔴")
            ent = row.get("Entreprise", "Inconnu")
            dt = row.get("Date", "N/A")
            st_name = row.get("Site", "N/A")
            admin_options_sidebar.append(f"{badge_icon} {sc:.0f}% | {ent} — {dt} ({st_name})")

        avg_score_val = (sum(scores_list) / len(scores_list)) if scores_list else 0.0
        
        st.sidebar.metric("Total d'audits enregistrés", total_audits_count)
        st.sidebar.metric("Conformité moyenne", f"{avg_score_val:.1f} %")
        st.sidebar.markdown("---")
        
        st.sidebar.markdown("### 📂 Sélection de l'Audit")
        selected_admin_sidebar_idx = st.sidebar.selectbox(
            "Consulter un audit spécifique :",
            range(len(admin_options_sidebar)),
            format_func=lambda x: admin_options_sidebar[x],
            key="sidebar_admin_audit_select"
        )
    else:
        st.sidebar.info("Aucun audit disponible dans la base.")

# ==========================================
# PAGE PRINCIPALE : FORMULAIRE PRESTATAIRE
# ==========================================
if app_mode == "📝 Formulaire Prestataire":
    st.markdown("""<div class="main-header">🛡️ Formulaire d'Audit Sécurité & HSE</div>""", unsafe_allow_html=True)
    st.markdown("""<div class="sub-header">Évaluation de conformité pour les entreprises extérieures. Merci de répondre à chaque question et de fournir obligatoirement une justification.</div>""", unsafe_allow_html=True)

    col_sc1, col_sc2, col_sc3 = st.columns([1, 1, 2])
    with col_sc1:
        st.metric("Score Global Oui", f"{global_score_pct:.1f} %")
    with col_sc2:
        st.metric("Questions renseignées", f"{answered_q_count} / {total_q}")
    with col_sc3:
        st.markdown(f"**Évaluation du score :** {get_color_badge(global_score_pct)}")
        st.progress(global_score_pct / 100.0)

    # --- REPRISE DE BROUILLON DYNAMIQUE ---
    with st.expander("📂 Reprendre un brouillon enregistré auparavant (Optionnel)", expanded=False):
        uploaded_draft = st.file_uploader("Importez votre fichier de brouillon (.json) :", type=["json"], key="draft_importer")
        if uploaded_draft is not None:
            try:
                draft_data = json.load(uploaded_draft)
                
                if "Entreprise" in draft_data:
                    st.session_state["company_name_key"] = draft_data["Entreprise"]
                if "Déclarant" in draft_data:
                    st.session_state["auditor_name_key"] = draft_data["Déclarant"]
                if "Site" in draft_data:
                    st.session_state["site_location_key"] = draft_data["Site"]

                for q in QUESTIONS_DATA:
                    q_id = q["id"]
                    r_val = draft_data.get(f"{q_id}_Réponse")
                    j_val = draft_data.get(f"{q_id}_Justification")
                    if r_val in ["Oui", "Non", "N/A"]:
                        st.session_state[f"status_{q_id}"] = r_val
                    if j_val is not None:
                        st.session_state[f"justif_{q_id}"] = j_val

                st.success("✅ Brouillon chargé avec succès ! Toutes les réponses et justifications ont été réinjectées.")
            except Exception as e:
                st.error(f"⚠️ Fichier de brouillon invalide : {e}")

    st.subheader("1. Informations de l'Entreprise Extérieure")
    col1, col2 = st.columns(2)
    with col1:
        company_name = st.text_input("Nom de l'entreprise *", key="company_name_key", placeholder="Ex: ABC Construction")
        auditor_name = st.text_input("Nom du déclarant / Représentant HSE *", key="auditor_name_key", placeholder="Ex: Jean Dupont")
    with col2:
        site_location = st.text_input("Site / Chantier concerné *", key="site_location_key", placeholder="Ex: Usine Amiens - Zone B")
        audit_date = st.date_input("Date de soumission", value=datetime.today())

    st.markdown("---")
    st.subheader("2. Grille d'Évaluation des Exigences Sécurité")

    if selected_theme_nav == "📋 Tous les thèmes":
        filtered_themes = THEME_LIST
    else:
        filtered_themes = [selected_theme_nav]

    responses = {}
    uploaded_files_dict = {}

    for t_name in filtered_themes:
        st.markdown(f"## 📌 {t_name}")
        theme_questions = [q for q in QUESTIONS_DATA if q["theme"] == t_name]
        
        for q in theme_questions:
            st.markdown(
                f'<div class="question-card"><div class="question-title-big">[{q["id"]}] {q["q"]}</div></div>',
                unsafe_allow_html=True
            )
            
            if q["g"]:
                st.markdown(f'<div class="guidance-box"><b>Attentes & Guidance :</b><br>{q["g"]}</div>', unsafe_allow_html=True)

            c1, c2 = st.columns([1, 2])
            with c1:
                status = st.radio(
                    f"Réponse pour [{q['id']}]",
                    options=["Oui", "Non", "N/A"],
                    index=None,
                    horizontal=True,
                    key=f"status_{q['id']}"
                )
            with c2:
                justification = st.text_area(
                    f"Justification obligatoire pour [{q['id']}] *",
                    placeholder="Justification OBLIGATOIRE (procédure interne, preuve, précision ou plan d'action)...",
                    key=f"justif_{q['id']}",
                    height=90
                )
            
            file_attached = st.file_uploader(
                f"📎 Pièce justificative pour [{q['id']}] (Optionnel)",
                type=["pdf", "png", "jpg", "jpeg", "docx", "xlsx"],
                key=f"file_{q['id']}"
            )
            
            responses[q['id']] = {
                "status": status,
                "justification": justification
            }
            if file_attached is not None:
                uploaded_files_dict[q['id']] = file_attached
                
            st.markdown("<hr style='margin: 20px 0; border-top: 2px dashed #94A3B8;'>", unsafe_allow_html=True)

    for q in QUESTIONS_DATA:
        if q['id'] not in responses:
            q_stat = st.session_state.get(f"status_{q['id']}", None)
            q_just = st.session_state.get(f"justif_{q['id']}", "")
            responses[q['id']] = {
                "status": q_stat,
                "justification": q_just
            }

    submit_button = st.button("🚀 VALIDER ET ENVOYER L'AUDIT")

    if submit_button:
        missing_fields = []
        if not company_name: missing_fields.append("Nom de l'entreprise")
        if not auditor_name: missing_fields.append("Nom du déclarant")
        if not site_location: missing_fields.append("Site / Chantier")
        
        unanswered = [q_id for q_id, res in responses.items() if res["status"] is None]
        unjustified = [q_id for q_id, res in responses.items() if res["status"] is not None and not res["justification"].strip()]
                
        if missing_fields:
            st.error(f"⚠️ Veuillez remplir les informations obligatoires : {', '.join(missing_fields)}.")
        elif unanswered:
            st.warning(f"⚠️ Vous devez répondre à toutes les questions ({len(QUESTIONS_DATA)} au total). Question(s) non renseignée(s) : {', '.join(unanswered)}")
        elif unjustified:
            st.warning(f"⚠️ La justification est OBLIGATOIRE pour chaque question (même pour les réponses 'Oui'). Question(s) sans justification : {', '.join(unjustified)}")
        else:
            st.info("⏳ Enregistrement de l'audit en cours...")
            
            record = {
                "Date": str(audit_date),
                "Entreprise": company_name,
                "Déclarant": auditor_name,
                "Site": site_location,
            }
            
            for q_id, res in responses.items():
                record[f"{q_id}_Réponse"] = res["status"]
                record[f"{q_id}_Justification"] = res["justification"]
                
                if q_id in uploaded_files_dict:
                    file_obj = uploaded_files_dict[q_id]
                    file_bytes = file_obj.getvalue()
                    file_b64 = base64.b64encode(file_bytes).decode("utf-8")
                    
                    record[f"{q_id}_Fichier"] = {
                        "name": file_obj.name,
                        "type": file_obj.type,
                        "data": file_b64
                    }
                else:
                    record[f"{q_id}_Fichier"] = None
            
            # 1. Sauvegarde locale
            enregistrer_audit_fichier_local(record)
            st.session_state["local_audits"].append(record)
            
            # 2. Synchronisation Airtable
            succes_airtable = enregistrer_audit_airtable(record)
            if succes_airtable:
                st.success("✅ Audit enregistré avec succès dans la base distante Airtable !")
            else:
                st.success("✅ Audit enregistré localement dans la session !")

            st.balloons()
            
            # --- TÉLÉCHARGEMENT DU RAPPORT EXCEL MIS EN FORME POUR LE PRESTATAIRE ---
            st.markdown("### 📥 Télécharger votre rapport d'audit")
            excel_bytes_user = generer_excel_formatted(record)
            nom_entreprise_clean = str(company_name).replace(" ", "_")
            
            st.download_button(
                label="📥 TÉLÉCHARGER MON COMPTE-RENDU EXCEL MIS EN FORME (.XLSX)",
                data=excel_bytes_user,
                file_name=f"Audit_{nom_entreprise_clean}_{audit_date}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary"
            )

    st.markdown("---")
    st.subheader("💾 Sauvegarder votre avancement")
    draft_export = {
        "Entreprise": company_name if 'company_name' in locals() else "",
        "Déclarant": auditor_name if 'auditor_name' in locals() else "",
        "Site": site_location if 'site_location' in locals() else "",
    }
    if 'responses' in locals():
        for q_id, res in responses.items():
            draft_export[f"{q_id}_Réponse"] = res["status"] if res["status"] is not None else ""
            draft_export[f"{q_id}_Justification"] = res["justification"]

    st.download_button(
        label="💾 Télécharger le fichier de brouillon (.json)",
        data=json.dumps(draft_export, ensure_ascii=False, indent=2),
        file_name=f"Brouillon_Audit_{company_name if 'company_name' in locals() and company_name else 'Incomplet'}.json",
        mime="application/json"
    )

# ==========================================
# PAGE PRINCIPALE : ESPACE ADMINISTRATEUR HSE
# ==========================================
else:
    st.markdown("""<div class="main-header">🔒 Espace d'Administration HSE</div>""", unsafe_allow_html=True)
    
    ADMIN_PASSWORD = st.secrets.get("admin_password", "HSE2026Securite!")
    input_pwd = st.text_input("🔑 Saisissez le mot de passe Administrateur :", type="password")
    
    if input_pwd == "":
        st.info("🔒 Cet espace est strictement réservé à la consultation administrateur.")
    elif input_pwd != ADMIN_PASSWORD:
        st.error("❌ Mot de passe incorrect.")
    else:
        st.success("🔓 Accès administrateur autorisé.")
        st.markdown("""<div class="sub-header">Consultez l'ensemble des audits validés, analysez la traçabilité globale et exportez les rapports Excel.</div>""", unsafe_allow_html=True)
        
        # --- SAUVEGARDE & RESTAURATION MANUELLE DE LA BASE ---
        with st.expander("🛠️ Gestion de la Sauvegarde / Restauration de la Base de Données", expanded=False):
            st.markdown("Utilisez ces outils pour exporter ou réimporter manuellement vos sauvegardes.")
            col_b1, col_b2 = st.columns(2)
            
            with col_b1:
                df_curr = charger_tous_les_audits()
                json_backup = df_curr.to_json(orient="records", force_ascii=False, indent=2) if not df_curr.empty else "[]"
                st.download_button(
                    label="💾 Exporter la base d'audits (.json)",
                    data=json_backup,
                    file_name=f"Sauvegarde_Base_Audits_{datetime.now().strftime('%Y%m%d_%H%M')}.json",
                    mime="application/json"
                )
            
            with col_b2:
                uploaded_db = st.file_uploader("📥 Importer/Restaurer une sauvegarde (.json)", type=["json"], key="restore_db_uploader")
                if uploaded_db is not None:
                    try:
                        imported_audits = json.load(uploaded_db)
                        if isinstance(imported_audits, list):
                            st.session_state["local_audits"] = imported_audits
                            with open(DB_FILE, "w", encoding="utf-8") as f:
                                json.dump(imported_audits, f, ensure_ascii=False, indent=2)
                            st.success(f"✅ Base restaurée avec succès ({len(imported_audits)} audits récupérés) !")
                            st.cache_data.clear()
                    except Exception as err:
                        st.error(f"Erreur lors de l'importation : {err}")

        if st.button("🔄 Actualiser la liste des audits"):
            st.cache_data.clear()

        df_audits = charger_tous_les_audits()

        if df_audits.empty:
            st.warning("⚠️ Aucun audit n'a encore été enregistré.")
        else:
            # --- TABLEAU RÉCAPITULATIF DE TRAÇABILITÉ GLOBALE ---
            st.markdown("---")
            st.subheader("📋 Tableau de Bord & Traçabilité des Audits")
            
            recap_data = []
            for idx, row in df_audits.iterrows():
                oui_c = sum(1 for col in row.index if str(col).endswith("_Réponse") and str(row[col]).strip() == "Oui")
                non_c = sum(1 for col in row.index if str(col).endswith("_Réponse") and str(row[col]).strip() == "Non")
                na_c = sum(1 for col in row.index if str(col).endswith("_Réponse") and str(row[col]).strip() == "N/A")
                t_app = oui_c + non_c
                score = (oui_c / t_app * 100) if t_app > 0 else 0.0
                statut_str = "🟢 Conforme" if score >= 80 else ("🟠 À améliorer" if score >= 50 else "🔴 Non conforme")

                recap_data.append({
                    "Date": row.get("Date", "N/A"),
                    "Entreprise": row.get("Entreprise", "N/A"),
                    "Site": row.get("Site", "N/A"),
                    "Déclarant": row.get("Déclarant", "N/A"),
                    "Score Global (%)": f"{score:.1f} %",
                    "Statut": statut_str,
                    "Oui": oui_c,
                    "Non": non_c,
                    "N/A": na_c
                })

            df_recap = pd.DataFrame(recap_data)
            st.dataframe(df_recap, use_container_width=True, hide_index=True)

            # --- SELECTION DE LA FICHE DETAIL ---
            selected_idx = st.session_state.get("sidebar_admin_audit_select", 0)
            if selected_idx >= len(df_audits):
                selected_idx = 0
                
            selected_row = df_audits.iloc[selected_idx]
            
            st.markdown("---")
            st.subheader(f"📄 Fiche Détail Audit : {selected_row.get('Entreprise', 'N/A')}")
            
            c1, c2, c3, c4 = st.columns(4)
            c1.metric("🏢 Entreprise", str(selected_row.get("Entreprise", "N/A")))
            c2.metric("👤 Déclarant", str(selected_row.get("Déclarant", "N/A")))
            c3.metric("📍 Site", str(selected_row.get("Site", "N/A")))
            c4.metric("📅 Date", str(selected_row.get("Date", "N/A")))
            
            oui_count = sum(1 for col in selected_row.index if str(col).endswith("_Réponse") and str(selected_row[col]).strip() == "Oui")
            non_count = sum(1 for col in selected_row.index if str(col).endswith("_Réponse") and str(selected_row[col]).strip() == "Non")
            na_count = sum(1 for col in selected_row.index if str(col).endswith("_Réponse") and str(selected_row[col]).strip() == "N/A")
            
            total_applicable_admin = oui_count + non_count
            score_percentage_admin = (oui_count / total_applicable_admin * 100) if total_applicable_admin > 0 else 0.0
            
            st.markdown("### 📊 Score Total & Conformité Globale")
            m1, m2, m3, m4 = st.columns(4)
            m1.metric("Score Global Oui (%)", f"{score_percentage_admin:.1f} %")
            m2.metric("✅ Conforme (Oui)", oui_count)
            m3.metric("❌ Non Conforme (Non)", non_count)
            m4.metric("⚪ Non Applicable (N/A)", na_count)
            
            st.markdown(f"**Évaluation du score total :** {get_color_badge(score_percentage_admin)}")
            st.progress(score_percentage_admin / 100.0)
            
            # Export Excel
            st.markdown(" ")
            excel_bytes = generer_excel_formatted(selected_row)
            nom_entreprise_clean = str(selected_row.get("Entreprise", "Audit")).replace(" ", "_")
            date_clean = str(selected_row.get("Date", "2026"))

            st.download_button(
                label="📥 TÉLÉCHARGER LE RAPPORT EXCEL MIS EN FORME (.XLSX)",
                data=excel_bytes,
                file_name=f"Audit_{nom_entreprise_clean}_{date_clean}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary"
            )
            
            st.markdown("---")
            st.markdown("### 🎯 Score Détaillé par Thème (5 Thèmes Sécurité)")
            
            cols_theme = st.columns(5)
            
            for idx_t, theme_name in enumerate(THEME_LIST):
                theme_questions = [q for q in QUESTIONS_DATA if q["theme"] == theme_name]
                t_oui = 0
                t_non = 0
                t_na = 0
                
                for q in theme_questions:
                    res_val = str(selected_row.get(f"{q['id']}_Réponse", "")).strip()
                    if res_val == "Oui":
                        t_oui += 1
                    elif res_val == "Non":
                        t_non += 1
                    elif res_val == "N/A":
                        t_na += 1
                
                t_appl = t_oui + t_non
                t_score = (t_oui / t_appl * 100) if t_appl > 0 else 0.0
                
                with cols_theme[idx_t]:
                    st.markdown(f"**Thème {idx_t+1}**")
                    st.caption(f"*{theme_name}*")
                    st.metric("Score", f"{t_score:.0f} %", f"{t_oui}/{t_appl} Oui")
                    
                    if t_score >= 80:
                        st.success("🟢 Conforme")
                    elif t_score >= 50:
                        st.warning("🟠 Moyen")
                    else:
                        st.error("🔴 Critique")
                        
                    st.caption(f"Oui: {t_oui} | Non: {t_non} | N/A: {t_na}")
            
            st.markdown("---")
            st.markdown("### 🔍 Détail des Questions par Thème")
            
            for theme_name in THEME_LIST:
                st.markdown(f"#### 📌 {theme_name}")
                theme_q = [q for q in QUESTIONS_DATA if q["theme"] == theme_name]
                
                for q in theme_q:
                    q_id = q["id"]
                    resp_val = str(selected_row.get(f"{q_id}_Réponse", "Non renseigné"))
                    justif_val = str(selected_row.get(f"{q_id}_Justification", "Aucune justification"))
                    file_val = selected_row.get(f"{q_id}_Fichier")
                    
                    if resp_val == "Oui":
                        badge = "🟢 **Oui**"
                    elif resp_val == "Non":
                        badge = "🔴 **Non**"
                    else:
                        badge = "⚪ **N/A**"
                    
                    with st.expander(f"[{q_id}] {q['q']} — {badge}"):
                        st.write(f"**Exigence / Question :** {q['q']}")
                        if q['g']:
                            st.caption(f"**Attentes :** {q['g']}")
                        st.write(f"**Statut :** {badge}")
                        st.write(f"**Justification entreprise :** {justif_val}")
                        
                        if isinstance(file_val, dict) and file_val.get("data"):
                            file_bytes = base64.b64decode(file_val["data"])
                            st.download_button(
                                label=f"📎 Télécharger {file_val.get('name', 'pièce jointe')}",
                                data=file_bytes,
                                file_name=file_val.get("name", "document"),
                                mime=file_val.get("type", "application/octet-stream"),
                                key=f"dl_admin_{selected_idx}_{q_id}"
                            )
                        elif isinstance(file_val, str) and file_val.startswith("http"):
                            st.markdown(f"📎 **Pièce jointe :** [Ouvrir le document]({file_val})")
                        else:
                            st.caption("📎 Aucune pièce jointe transmise pour cette question.")
